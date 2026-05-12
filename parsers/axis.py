"""
Axis Parser
-----------
File:     Encrypted .xlsx inside an unprotected zip
Password: On the file itself (CDFV2 encryption) — NOT on the zip
Strategy: FundStrategyCode column (col 5) — also in the filename
Sheet:    'Holding Report1'
Layout:   Row 1 = column headers, data from row 2, no footer rows
Quantity: NetBalance (col 26 = logical), SaleStockQty (col 25 = saleable)
"""
import logging
import io
import openpyxl
import msoffcrypto
from datetime import datetime
from .base import BaseParser, HoldingRecord, ParseResult

logger = logging.getLogger(__name__)

# Column indices (0-based) — 30 columns total
COL_SR_NO        = 0
COL_UCC          = 1   # client code
COL_CLIENT_NAME  = 2
COL_CLIENT_ID    = 3
COL_PARENT_NAME  = 4
COL_STRATEGY     = 5   # FundStrategyCode — broker strategy code
COL_SECURITY     = 6   # SecurityName
COL_ISIN         = 7
COL_ISIN_SUBTYPE = 8
COL_PRODUCT_TYPE = 9
COL_INVESTOR     = 10
# ... pending/physical columns 11-20 ...
COL_DEMAT_FREE   = 21
COL_DEMAT_LOCKED = 22
COL_DEMAT_PLEDGE = 23
COL_DEMAT_FREEZ  = 24
COL_SALE_STOCK   = 25  # saleable quantity
COL_NET_BALANCE  = 26  # logical holding
COL_LTP_EXCH     = 27
COL_LTP_DATE     = 28
COL_LTP_PRICE    = 29


class AxisParser(BaseParser):
    source_name = 'axis'

    def parse_file(self, file_path: str, date_str: str,
                   file_password: str = '') -> ParseResult:
        result = ParseResult(source=self.source_name, file_path=file_path)
        holding_date = self.format_date(date_str)

        try:
            wb = self._open_workbook(file_path, file_password)

            sheet_name = 'Holding Report1'
            if sheet_name not in wb.sheetnames:
                result.error = f"Expected sheet '{sheet_name}' not found. Sheets: {wb.sheetnames}"
                return result

            ws = wb[sheet_name]
            broker_codes = set()

            # Row 1 is headers, data starts from row 2
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or row[COL_SR_NO] is None:
                    continue

                client_code   = str(row[COL_UCC] or '').strip()
                broker_code   = str(row[COL_STRATEGY] or '').strip()
                security_name = str(row[COL_SECURITY] or '').strip()
                isin          = str(row[COL_ISIN] or '').strip()

                # Strip " - INR" suffix from security names (Axis appends this)
                security_name = security_name.replace(' - INR', '').strip()

                if not isin or not broker_code or len(isin) < 5:
                    continue

                logical  = self.clean_number(row[COL_NET_BALANCE])
                saleable = self.clean_number(row[COL_SALE_STOCK])

                # Use ISIN as security code since Axis has no short code column
                # The instrument code will be populated from security name
                security_code = isin

                broker_codes.add(broker_code)
                result.records.append(HoldingRecord(
                    broker_code      = broker_code,
                    client_id        = client_code,
                    security_code    = security_code,
                    security_name    = security_name,
                    isin             = isin,
                    logical_holding  = logical,
                    saleable_holding = saleable,
                    holding_date     = holding_date,
                    source           = self.source_name,
                ))

            result.broker_codes_found = sorted(broker_codes)
            logger.info(f"Axis: parsed {result.record_count} records, "
                        f"broker codes: {result.broker_codes_found}")

        except Exception as e:
            result.error = str(e)
            logger.error(f"Axis parse error: {e}", exc_info=True)

        return result

    def _open_workbook(self, file_path: str, password: str):
        """Open xlsx — handles both plain and CDFV2-encrypted files.

        Raises a parser-friendly exception with the specific failure
        cause on failure (wrong password / corrupt file / msoffcrypto
        bug). Previous behaviour returned None and surfaced as a
        generic "Could not open Axis file" — the actual error was only
        in the log file, which operators rarely check.
        """
        plain_error = None

        # Try plain open first (file may not be encrypted)
        try:
            return openpyxl.load_workbook(file_path)
        except Exception as e:
            plain_error = e

        # Encrypted: need a password
        if not password:
            raise RuntimeError(
                f"Axis file is encrypted but no file_password is configured "
                f"in Settings (sources.json[axis].file_password). "
                f"Plain-open error: {plain_error}"
            )

        try:
            with open(file_path, 'rb') as f:
                office_file = msoffcrypto.OfficeFile(f)
                office_file.load_key(password=password)
                decrypted = io.BytesIO()
                office_file.decrypt(decrypted)
                decrypted.seek(0)
                return openpyxl.load_workbook(decrypted)
        except msoffcrypto.exceptions.InvalidKeyError as e:
            raise RuntimeError(
                f"Axis file decryption failed: wrong password. Check "
                f"sources.json[axis].file_password against the password "
                f"on the latest Axis email. Underlying: {e}"
            )
        except Exception as e:
            raise RuntimeError(
                f"Axis file decryption failed (msoffcrypto: {type(e).__name__}: {e}; "
                f"plain-open error: {plain_error})"
            )
