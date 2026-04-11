"""
ICICI Parser
------------
File:     .xlsx inside a password-protected zip
Sheet:    RPT_EndClientHolding
Password: On the zip (not the file)
Strategy: 'Scheme' column (col index 2) — e.g. GOLDWEMPMS, GOLDWEVPMS, GOLDCGT
Quantity: Saleable Position (col 10), Pending Buy (col 8), Pending Sell (col 9)
Header:   Row 1 = title, Row 2 = client info, Row 3 = column headers
Footer:   Last 2 rows (date row + disclaimer row)
"""
import logging
import openpyxl
from datetime import datetime
from .base import BaseParser, HoldingRecord, ParseResult

logger = logging.getLogger(__name__)

# Column indices (0-based) in RPT_EndClientHolding
COL_CLIENT_CODE   = 0
COL_CLIENT_NAME   = 1
COL_SCHEME        = 2   # broker strategy code
COL_INSTR_NAME    = 3
COL_INSTR_SUBTYPE = 4
COL_ISIN          = 5
COL_CLIENT_INSTR  = 6   # usually blank
COL_INSTR_CODE    = 7   # broker security code
COL_PEND_BUY      = 8
COL_PEND_SELL     = 9
COL_SALEABLE      = 10
COL_LOCATION      = 11
COL_BLOCK_TYPE    = 12


class ICICIParser(BaseParser):
    source_name = 'icici'

    def parse_file(self, file_path: str, date_str: str,
                   file_password: str = '') -> ParseResult:
        result = ParseResult(source=self.source_name, file_path=file_path)
        holding_date = self.format_date(date_str)

        try:
            wb = openpyxl.load_workbook(file_path)

            if 'RPT_EndClientHolding' not in wb.sheetnames:
                result.error = f"Expected sheet 'RPT_EndClientHolding' not found. Sheets: {wb.sheetnames}"
                return result

            ws = wb['RPT_EndClientHolding']
            all_rows = list(ws.iter_rows(values_only=True))
            wb.close()

            # Rows 1-2: title + client info (skip)
            # Row 3: column headers (skip)
            # Data starts from row 4 (index 3)
            # Last 2 rows: date row + disclaimer row (skip)
            data_rows = all_rows[3:-2]

            broker_codes = set()
            for row in data_rows:
                if not row or row[COL_CLIENT_CODE] is None:
                    continue
                # Skip any row where client code looks like a date/disclaimer
                client_code = str(row[COL_CLIENT_CODE]).strip()
                if client_code in ('', 'Date:', 'Disclaimer:'):
                    continue
                if isinstance(row[COL_CLIENT_CODE], datetime):
                    continue

                broker_code = str(row[COL_SCHEME] or '').strip()
                isin        = str(row[COL_ISIN] or '').strip()
                instr_code  = str(row[COL_INSTR_CODE] or '').strip()
                instr_name  = str(row[COL_INSTR_NAME] or '').strip()

                if not isin or not broker_code:
                    continue

                saleable   = self.clean_number(row[COL_SALEABLE])
                pend_buy   = self.clean_number(row[COL_PEND_BUY])
                pend_sell  = self.clean_number(row[COL_PEND_SELL])
                # Logical = Saleable + Pending Buy - Pending Sell
                logical    = saleable + pend_buy - pend_sell

                broker_codes.add(broker_code)
                result.records.append(HoldingRecord(
                    broker_code      = broker_code,
                    client_id        = client_code,
                    security_code    = instr_code,
                    security_name    = instr_name,
                    isin             = isin,
                    logical_holding  = logical,
                    saleable_holding = saleable,
                    holding_date     = holding_date,
                    source           = self.source_name,
                ))

            result.broker_codes_found = sorted(broker_codes)
            logger.info(f"ICICI: parsed {result.record_count} records, "
                        f"broker codes: {result.broker_codes_found}")

        except Exception as e:
            result.error = str(e)
            logger.error(f"ICICI parse error: {e}", exc_info=True)

        return result
